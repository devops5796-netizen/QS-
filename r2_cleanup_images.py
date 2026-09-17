"""
QatarSale R2 cleanup script
============================
- Walks ALL dates under a base prefix (default: qatarsale/) on Cloudflare R2.
- For every category found under each date, it:
    1. Deletes everything inside that category's "images/" subfolder.
    2. Opens every Excel file under "excel/" (every sheet) and drops any
       column whose header matches one of the known image-reference columns.
    3. Opens every JSON file under "json/" and removes those same keys
       from every record.

Known image-reference column/key names (case-insensitive, any category):
    - images_local_paths
    - image_r2_key
    - r2_image

Run with --dry-run first to see exactly what would be touched before it
deletes/modifies anything for real.

Requirements:
    pip install boto3 openpyxl

Environment variables expected (fill these in / export them before running):
    R2_ACCOUNT_ID
    R2_ACCESS_KEY_ID
    R2_SECRET_ACCESS_KEY
    R2_BUCKET_NAME
"""

import argparse
import io
import json
import os
import re
import sys

import boto3
from botocore.config import Config
import openpyxl

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

BASE_PREFIX = "qatarsale/"           # e.g. qatarsale/year=2026/month=09/day=01/
TARGET_COLUMNS = {"images_local_paths", "image_r2_key", "r2_image"}

DATE_PREFIX_RE = re.compile(r"^year=\d{4}/month=\d{2}/day=\d{2}/$")


def get_client():
    account_id = os.environ["R2_ACCOUNT_ID"]
    return boto3.client(
        "s3",
        endpoint_url=f"https://{account_id}.r2.cloudflarestorage.com",
        aws_access_key_id=os.environ["R2_ACCESS_KEY_ID"],
        aws_secret_access_key=os.environ["R2_SECRET_ACCESS_KEY"],
        config=Config(signature_version="s3v4"),
        region_name="auto",
    )


def list_common_prefixes(s3, bucket, prefix):
    """One level of 'folders' under `prefix` (uses delimiter='/')."""
    paginator = s3.get_paginator("list_objects_v2")
    prefixes = []
    for page in paginator.paginate(Bucket=bucket, Prefix=prefix, Delimiter="/"):
        for cp in page.get("CommonPrefixes", []):
            prefixes.append(cp["Prefix"])
    return prefixes


def list_all_date_prefixes(s3, bucket):
    """qatarsale/year=YYYY/month=MM/day=DD/ for every date that exists."""
    date_prefixes = []
    for year_prefix in list_common_prefixes(s3, bucket, BASE_PREFIX):
        for month_prefix in list_common_prefixes(s3, bucket, year_prefix):
            for day_prefix in list_common_prefixes(s3, bucket, month_prefix):
                date_prefixes.append(day_prefix)
    return sorted(date_prefixes)


def list_all_objects(s3, bucket, prefix):
    paginator = s3.get_paginator("list_objects_v2")
    keys = []
    for page in paginator.paginate(Bucket=bucket, Prefix=prefix):
        for obj in page.get("Contents", []):
            keys.append(obj["Key"])
    return keys


def delete_keys(s3, bucket, keys, dry_run):
    if not keys:
        return
    print(f"    - deleting {len(keys)} object(s) under images/")
    if dry_run:
        return
    # S3/R2 batch delete takes max 1000 keys per call
    for i in range(0, len(keys), 1000):
        chunk = keys[i:i + 1000]
        s3.delete_objects(
            Bucket=bucket,
            Delete={"Objects": [{"Key": k} for k in chunk]},
        )


def clean_excel_object(s3, bucket, key, dry_run):
    obj = s3.get_object(Bucket=bucket, Key=key)
    data = obj["Body"].read()
    wb = openpyxl.load_workbook(io.BytesIO(data))

    changed = False
    for ws in wb.worksheets:
        if ws.max_row == 0:
            continue
        header_row = 1
        headers = {}
        for col_idx, cell in enumerate(ws[header_row], start=1):
            if cell.value is not None:
                headers[str(cell.value).strip().lower()] = col_idx

        cols_to_delete = sorted(
            (idx for name, idx in headers.items() if name in TARGET_COLUMNS),
            reverse=True,
        )
        if cols_to_delete:
            changed = True
            found_names = [n for n in headers if headers[n] in cols_to_delete]
            print(f"    - {key} [sheet '{ws.title}']: dropping columns {found_names}")
            if not dry_run:
                for col_idx in cols_to_delete:
                    ws.delete_cols(col_idx, 1)

    if changed and not dry_run:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        s3.put_object(Bucket=bucket, Key=key, Body=buf.getvalue())
    return changed


def strip_keys_from_json_obj(obj, keys_lower):
    """Recursively drop matching keys from dicts, whether top-level dict,
    list of dicts, or a dict wrapping a list of dicts."""
    if isinstance(obj, dict):
        return {
            k: strip_keys_from_json_obj(v, keys_lower)
            for k, v in obj.items()
            if k.strip().lower() not in keys_lower
        }
    if isinstance(obj, list):
        return [strip_keys_from_json_obj(item, keys_lower) for item in obj]
    return obj


def clean_json_object(s3, bucket, key, dry_run):
    obj = s3.get_object(Bucket=bucket, Key=key)
    raw = obj["Body"].read()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        print(f"    ! skipping {key}: not valid JSON")
        return False

    original_str = json.dumps(data, sort_keys=True, ensure_ascii=False)
    cleaned = strip_keys_from_json_obj(data, TARGET_COLUMNS)
    cleaned_str = json.dumps(cleaned, sort_keys=True, ensure_ascii=False)

    changed = original_str != cleaned_str
    if changed:
        print(f"    - {key}: removed matching key(s)")
        if not dry_run:
            s3.put_object(
                Bucket=bucket,
                Key=key,
                Body=json.dumps(cleaned, ensure_ascii=False, indent=2).encode("utf-8"),
            )
    return changed


def process_category(s3, bucket, category_prefix, dry_run):
    category_name = category_prefix.rstrip("/").split("/")[-1]
    print(f"  Category: {category_name}")

    # 1) delete images/ subfolder entirely
    images_prefix = category_prefix + "images/"
    image_keys = list_all_objects(s3, bucket, images_prefix)
    delete_keys(s3, bucket, image_keys, dry_run)

    # 2) clean excel/ files (any number of files, any number of sheets)
    excel_prefix = category_prefix + "excel/"
    for key in list_all_objects(s3, bucket, excel_prefix):
        if key.lower().endswith((".xlsx", ".xlsm")):
            clean_excel_object(s3, bucket, key, dry_run)

    # 3) clean json/ files (any number of files)
    json_prefix = category_prefix + "json/"
    for key in list_all_objects(s3, bucket, json_prefix):
        if key.lower().endswith(".json"):
            clean_json_object(s3, bucket, key, dry_run)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true",
                         help="Show what would be deleted/changed without touching R2")
    parser.add_argument("--date", default=None,
                         help="Only process one date, e.g. year=2026/month=09/day=01/")
    args = parser.parse_args()

    bucket = os.environ["R2_BUCKET_NAME"]
    s3 = get_client()

    if args.date:
        date_prefixes = [BASE_PREFIX + args.date.rstrip("/") + "/"]
    else:
        date_prefixes = list_all_date_prefixes(s3, bucket)

    print(f"Found {len(date_prefixes)} date(s) to process."
          f"{' [DRY RUN]' if args.dry_run else ''}")

    for date_prefix in date_prefixes:
        print(f"\n{'=' * 70}\n📅 {date_prefix}\n{'=' * 70}")
        category_prefixes = list_common_prefixes(s3, bucket, date_prefix)
        for category_prefix in category_prefixes:
            process_category(s3, bucket, category_prefix, args.dry_run)

    print("\nDone." + (" (dry run — nothing was actually changed)" if args.dry_run else ""))


if __name__ == "__main__":
    sys.exit(main())